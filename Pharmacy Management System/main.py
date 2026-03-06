from tkinter import *
from tkinter import messagebox
from tkinter.ttk import Treeview,Combobox,Style
from PIL import ImageTk,Image
from tkcalendar import DateEntry
import sqlite3
import datetime
connect=sqlite3.connect('pharmacy.db')
cursor=connect.cursor()
cursor.execute("""
create table if not exists Medicine(
Id integer primary key autoincrement,
Name text,
Company_Name text,
Type text,
Usage text,
Mfg_Date date,
Expiry_Date date,
Lot_No text,
Quantity integer,
Sold integer,
In_Storage integer,
Price integer
)
"""
)
cursor.execute("""
create table if not exists Medicine_Names(
Id integer primary key autoincrement,
Name text,
Total_Quantity integer
)
"""
)
class Pharmacy:
    def __init__(self):
        self.Pharmacy_Image = ImageTk.PhotoImage(Image.open('pharmacy.jpg').resize((w, h)))
        self.Main_Background = Label(image=self.Pharmacy_Image)
        self.Main_Background.place(x=0, y=0)
        self.Title_Label = Label(text="Welcome To The Pharmacy", bg='#7FFFD4', fg='#0020C2',font=('Calibri', 40, 'bold'),bd=10, relief=RIDGE)
        self.Title_Label.pack(fill=X)
        self.Buttons_Frame = Frame(bd=10, relief=RIDGE, bg='#8FDFEA',padx=10,pady=10)
        self.Buttons_Frame.pack(pady=50)
        self.Add_Medicine_Button = Button(self.Buttons_Frame, text='Add Medicine',command=self.Add_Medicine, font=("Calibri", 15, "bold"), bg="green",fg='white', activebackground='green', activeforeground='white')
        self.Add_Medicine_Button.pack(pady=10)
        self.Show_Medicines_Button = Button(self.Buttons_Frame, text='Show Medicines',command=self.Show_Medicines, font=("Calibri", 15, "bold"), bg="green",fg='white', activebackground='green', activeforeground='white')
        self.Show_Medicines_Button.pack(pady=10)
        self.Show_Medicines_Quantity_Button = Button(self.Buttons_Frame, text='Show Medicines Quantity', command=self.Show_Medicines_Quantity,font=("Calibri", 15, "bold"), bg="green", fg='white', activebackground='green', activeforeground='white')
        self.Show_Medicines_Quantity_Button.pack(pady=10)
        self.Show_Expired_Medicines_Button = Button(self.Buttons_Frame, text='Show Expired Medicines', command=self.Show_Expired_Medicines,font=("Calibri", 15, "bold"), bg="green", fg='white',activebackground='green', activeforeground='white')
        self.Show_Expired_Medicines_Button.pack(pady=10)
        self.Show_Missing_Medicines_Button = Button(self.Buttons_Frame, text='Show Missing Medicines',command=self.Show_Missing_Medicines, font=("Calibri", 15, "bold"),bg="green", fg='white', activebackground='green',activeforeground='white')
        self.Show_Missing_Medicines_Button.pack(pady=10)
        self.Sell_Medicine_Button = Button(self.Buttons_Frame, text='Sell Medicine', command=self.Sell_Medicine,font=("Calibri", 15, "bold"), bg="green", fg='white',activebackground='green', activeforeground='white')
        self.Sell_Medicine_Button.pack(pady=10)
    def Remove_Main(self):
        self.Title_Label.pack_forget()
        self.Buttons_Frame.pack_forget()
    def Add_Medicine(self):
        self.Remove_Main()
        Add_Medicine()
    def Sell_Medicine(self):
        self.Remove_Main()
        Sell_Medicine()
    def Show_Medicines(self):
        self.Remove_Main()
        Show_Medicines()
    def Show_Medicines_Quantity(self):
        self.Remove_Main()
        Show_Medicines_Quantity()
    def Show_Expired_Medicines(self):
        self.Remove_Main()
        Show_Expired_Medicines()
    def Show_Missing_Medicines(self):
        self.Remove_Main()
        Show_Missing_Medicines()
class Add_Medicine:
    def __init__(self):
        self.Background_Image = ImageTk.PhotoImage(Image.open('medicine.jpg').resize((w, h)))
        self.Back_Image = ImageTk.PhotoImage(Image.open('back.png').resize((100, 100)))
        self.Main_Background = Label(image=self.Background_Image)
        self.Main_Background.place(x=0, y=0)
        self.Back_Add_Button = Button(image=self.Back_Image, command=self.Back_Button, bg='red', bd=0, activebackground='red')
        self.Back_Add_Button.place(x=0, y=0)
        self.Title_Label = Label(text="Add Medicine", bg='#7FFFD4', fg='#0020C2',font=('Calibri', 40, 'bold'), bd=10, relief=RIDGE)
        self.Title_Label.pack()
        self.Add_Medicine_Frame=Frame(bg='#8FDFEA',bd=10,relief=RIDGE,padx=20,pady=10)
        self.Add_Medicine_Frame.pack(pady=50)
        self.Medicine_Name_Label=Label(self.Add_Medicine_Frame, text="Medicine Name",font=('Calibri', 14,'bold'),bg='#8FDFEA')
        self.Medicine_Name_Label.grid(row=0, column=0)
        self.Medicine_Name_Entry = Entry(self.Add_Medicine_Frame,bd=2,font=('Calibri',15),relief=RIDGE)
        self.Medicine_Name_Entry.grid(row=0, column=1,pady=10)
        self.Company_Name_Label = Label(self.Add_Medicine_Frame, text="Company Name", font=('Calibri', 14, 'bold'),bg='#8FDFEA')
        self.Company_Name_Label.grid(row=1, column=0)
        self.Company_Name_Entry = Entry(self.Add_Medicine_Frame, bd=2, font=('Calibri', 15), relief=RIDGE)
        self.Company_Name_Entry.grid(row=1, column=1, pady=10)
        self.Usage_Label = Label(self.Add_Medicine_Frame, text="Usage", font=('Calibri', 14, 'bold'), bg='#8FDFEA')
        self.Usage_Label.grid(row=2, column=0)
        self.Usage_Entry = Entry(self.Add_Medicine_Frame, bd=2, font=('Calibri', 15), relief=RIDGE)
        self.Usage_Entry.grid(row=2, column=1, padx=30, pady=10)
        self.Type_Of_Medicine_Values=('Tablets','Capsules','Liquids','Drops','Spray','Creams','Inhalers','Injections','Patches')
        self.Type_Of_Medicine_Label = Label(self.Add_Medicine_Frame, text="Type Of Medicine", font=('Calibri', 14, 'bold'), bg='#8FDFEA')
        self.Type_Of_Medicine_Label.grid(row=3, column=0)
        self.Type_Of_Medicine_ComboBox = Combobox(self.Add_Medicine_Frame, font=('Calibri', 14), state="readonly", values=self.Type_Of_Medicine_Values)
        self.Type_Of_Medicine_ComboBox.grid(row=3, column=1, padx=20,pady=5)
        self.Mfg_Date_Label = Label(self.Add_Medicine_Frame, text="Mfg. Date", font=('Calibri', 14, 'bold'), bg='#8FDFEA')
        self.Mfg_Date_Label.grid(row=4, column=0)
        self.Mfg_Date_Entry = DateEntry(self.Add_Medicine_Frame, date_pattern = 'dd/MM/yyyy')
        self.Mfg_Date_Entry.grid(row=4, column=1,pady=5)
        self.Expiry_Date_Label = Label(self.Add_Medicine_Frame, text="Expiry Date",font=('Calibri', 14,'bold'),bg='#8FDFEA')
        self.Expiry_Date_Label.grid(row=5, column=0)
        self.Expiry_Date_Entry = DateEntry(self.Add_Medicine_Frame, date_pattern = 'dd/MM/yyyy')
        self.Expiry_Date_Entry.grid(row=5, column=1,pady=10)
        self.Lot_Num_Label=Label(self.Add_Medicine_Frame, text="Lot No.", font=('Calibri', 14, 'bold'), bg='#8FDFEA')
        self.Lot_Num_Label.grid(row=6, column=0, padx=5)
        self.Lot_Num_Entry=Entry(self.Add_Medicine_Frame, bd=2, font=('Calibri', 15), relief=RIDGE)
        self.Lot_Num_Entry.grid(row=6, column=1, padx=5,pady=10)
        self.Price_Label = Label(self.Add_Medicine_Frame, text="Price",font=('Calibri', 14,'bold'),bg='#8FDFEA')
        self.Price_Label.grid(row=7, column=0,padx=5)
        self.Price_Entry = Entry(self.Add_Medicine_Frame, bd=2, font=('Calibri', 15), relief=RIDGE)
        self.Price_Entry.grid(row=7, column=1,pady=10)
        self.Quantity_Label = Label(self.Add_Medicine_Frame, text="Quantity", font=('Calibri', 14, 'bold'), bg='#8FDFEA')
        self.Quantity_Label.grid(row=8, column=0)
        self.Quantity_Entry = Entry(self.Add_Medicine_Frame, bd=2, font=('Calibri', 15), relief=RIDGE)
        self.Quantity_Entry.grid(row=8, column=1, pady=10)
        self.Submit_Button=Button(self.Add_Medicine_Frame,text="Submit",command=self.Submit,font=('Calibri', 15), bg="green",fg='white', activebackground='green', activeforeground='white')
        self.Submit_Button.grid(row=9,column=0,columnspan=2)
        self.style=Style()
        self.style.theme_use("default")
        self.Submit_Button.bind('<Return>', self.Enter_Submit)
        self.Type_Of_Medicine_ComboBox.bind('<<ComboboxSelected>>', self.Type_Of_Medicine_Clicked)
    def Remove_Add_Medicine(self):
        self.Main_Background.place_forget()
        self.Back_Add_Button.place_forget()
        self.Title_Label.pack_forget()
        self.Add_Medicine_Frame.pack_forget()
    def Back_Button(self):
        self.Remove_Add_Medicine()
        Pharmacy()
    def Submit(self):
        today = datetime.date.today().strftime("%d/%m/%Y")
        diff=self.Expiry_Date_Entry.get_date()-self.Mfg_Date_Entry.get_date()
        if self.Medicine_Name_Entry.get()==''or self.Company_Name_Entry.get()=='' or self.Usage_Entry.get()=='' or self.Type_Of_Medicine_ComboBox.get()==''\
        or self.Lot_Num_Entry.get()=='' or self.Price_Entry.get()=='' or self.Quantity_Entry.get()=='':
            messagebox.showerror("Error","Some Fields Are Empty")
        elif self.Mfg_Date_Entry.get_date()>self.Expiry_Date_Entry.get_date():
            messagebox.showerror("Error","Mfg. Date Must Be Smaller Than Expiry Date")
        elif not(self.Type_Of_Medicine_ComboBox.get()=='Injections' or self.Type_Of_Medicine_ComboBox.get()=='Patches') and diff.days<730:
            messagebox.showerror("Error","The Difference Between Expiry Date And Mfg. Date Must Be At Least 2 Years")
        elif not self.Price_Entry.get().isdigit() and not self.Quantity_Entry.get().isdigit():
            messagebox.showerror("Error", "Wrong Entry In Price And Quantity")
        elif not self.Price_Entry.get().isdigit():
            messagebox.showerror("Error","Wrong Entry In Price")
        elif not self.Quantity_Entry.get().isdigit():
            messagebox.showerror("Error", "Wrong Entry In Quantity")
        else:
            if self.Type_Of_Medicine_ComboBox.get()=='Injections' or self.Type_Of_Medicine_ComboBox.get()=='Patches':
                Expiry_Date='-'
            else:
                Expiry_Date=self.Expiry_Date_Entry.get()
            cursor.execute(f"""
                insert into Medicine values(NULL,'{self.Medicine_Name_Entry.get()}','{self.Company_Name_Entry.get()}','{self.Type_Of_Medicine_ComboBox.get()}','{self.Usage_Entry.get()}'\
                ,'{self.Mfg_Date_Entry.get()}','{Expiry_Date}','{self.Lot_Num_Entry.get().upper()}',{int(self.Quantity_Entry.get())},0\
                ,{int(self.Quantity_Entry.get())},{int(self.Price_Entry.get())}
                )""")
            connect.commit()
            cursor.execute(f"select Name from Medicine_Names where Name='{self.Medicine_Name_Entry.get()}'")
            res=cursor.fetchone()
            if not res:
                cursor.execute(f"insert into Medicine_Names values(NULL,'{self.Medicine_Name_Entry.get()}',{int(self.Quantity_Entry.get())})")
                connect.commit()
            else:
                cursor.execute(f"update Medicine_Names set Total_Quantity=Total_Quantity+{int(self.Quantity_Entry.get())} where Name='{self.Medicine_Name_Entry.get()}'")
                connect.commit()
            messagebox.showinfo("Add Medicine", "Medicine Added Successfully")
            self.Medicine_Name_Entry.delete(0,END)
            self.Company_Name_Entry.delete(0,END)
            self.Usage_Entry.delete(0,END)
            self.Type_Of_Medicine_ComboBox.set("")
            self.Lot_Num_Entry.delete(0,END)
            self.Price_Entry.delete(0,END)
            self.Quantity_Entry.delete(0,END)
            self.Mfg_Date_Entry.set_date(today)
            self.Expiry_Date_Entry.set_date(today)
    def Enter_Submit(self,e):
        self.Submit()
    def Type_Of_Medicine_Clicked(self,e):
        if self.Type_Of_Medicine_ComboBox.get()=='Injections' or self.Type_Of_Medicine_ComboBox.get()=='Patches':
            self.Expiry_Date_Label.grid_forget()
            self.Expiry_Date_Entry.grid_forget()
        else:
            self.Expiry_Date_Label.grid(row=5, column=0)
            self.Expiry_Date_Entry.grid(row=5, column=1, pady=10)
class Sell_Medicine:
    def __init__(self):
        self.Background_Image = ImageTk.PhotoImage(Image.open('medicine.jpg').resize((w, h)))
        self.Back_Image = ImageTk.PhotoImage(Image.open('back.png').resize((100, 100)))
        self.Main_Background = Label(image=self.Background_Image)
        self.Main_Background.place(x=0, y=0)
        self.Back_Add_Button = Button(image=self.Back_Image, command=self.Back_Button, bg='red', bd=0,activebackground='red')
        self.Back_Add_Button.place(x=0, y=0)
        self.Title_Label = Label(text="Sell Medicine", bg='#7FFFD4', fg='#0020C2', font=('Calibri', 40, 'bold'), bd=10,relief=RIDGE)
        self.Title_Label.pack()
        self.Results_Frame=Frame(bd=10, relief=RIDGE,bg='#8FDFEA')
        self.Results_Frame.pack()
        self.Search_Frame = Frame(self.Results_Frame,bg='#8FDFEA',bd=10, relief=RIDGE, padx=10, pady=10)
        self.Search_Frame.pack(pady=10)
        self.Search_By_Label = Label(self.Search_Frame,bg='#8FDFEA', text="Search By", font=('Calibri', 14, 'bold'))
        self.Search_By_Label.grid(row=0, column=0)
        self.Search_Values = ("Medicine Name","Company Name","Type Of Medicine","Usage","Lot No.")
        self.Search_ComboBox = Combobox(self.Search_Frame, font=('Calibri', 14), state="readonly", values=self.Search_Values)
        self.Search_ComboBox.grid(row=0, column=1,padx=5)
        self.Search_Entry = Entry(self.Search_Frame, bd=2, font=('Calibri', 15), relief=RIDGE,state="disabled")
        self.Search_Entry.grid(row=0, column=2,padx=5)
        self.Search_Label = Label(self.Search_Entry, text="Search...", font=('Calibri', 13))
        self.Search_Label.place(x=10, y=0)
        self.Search_Button = Button(self.Search_Frame, command=self.Search_Button, text="Search", font=('Calibri', 14), bg="green", fg='white', activebackground='green', activeforeground='white')
        self.Search_Button.grid(row=0, column=3,padx=5)
        self.Show_All_Button = Button(self.Search_Frame, command=self.Show_All, text="Show All", font=('Calibri', 14), bg="green", fg='white', activebackground='green', activeforeground='white')
        self.Show_All_Button.grid(row=0, column=4, padx=10)
        self.Table_Frame = Frame(self.Results_Frame,bg='#8FDFEA',bd=10, relief=RIDGE)
        self.Table_Frame.pack(side=LEFT,pady=20)
        self.Columns = ("Medicine Name","Company Name","Type Of Medicine", "Usage", "Mfg. Date", "Expiry Date","Lot No.","In Storage","Price")
        self.ScrollBar_Y = Scrollbar(self.Table_Frame)
        self.ScrollBar_Y.pack(side=RIGHT, fill=Y)
        self.Search_Table = Treeview(self.Table_Frame, columns=self.Columns, show='headings',yscrollcommand=self.ScrollBar_Y.set)
        self.ScrollBar_Y.config(command=self.Search_Table.yview)
        self.Search_Table.column("Medicine Name", width=170)
        self.Search_Table.column("Company Name", width=170)
        self.Search_Table.column("Mfg. Date", width=90,anchor=CENTER)
        self.Search_Table.column("Expiry Date", width=90,anchor=CENTER)
        self.Search_Table.column("In Storage", width=90,anchor=CENTER)
        self.Search_Table.column("Lot No.", width=100)
        self.Search_Table.column("Usage", width=200)
        self.Search_Table.column("Type Of Medicine", width=130,anchor=CENTER)
        self.Search_Table.column("Price", width=60,anchor=CENTER)
        self.Search_Table.heading("Medicine Name", text="Medicine Name",anchor=CENTER)
        self.Search_Table.heading("Company Name", text="Company Name",anchor=CENTER)
        self.Search_Table.heading("Mfg. Date", text="Mfg. Date",anchor=CENTER)
        self.Search_Table.heading("Expiry Date", text="Expiry Date",anchor=CENTER)
        self.Search_Table.heading("Lot No.", text="Lot No.",anchor=CENTER)
        self.Search_Table.heading("In Storage", text="In Storage",anchor=CENTER)
        self.Search_Table.heading("Type Of Medicine", text="Type Of Medicine",anchor=CENTER)
        self.Search_Table.heading("Usage", text="Usage",anchor=CENTER)
        self.Search_Table.heading("Price", text="Price",anchor=CENTER)
        self.Search_Table.pack()
        self.Sell_Medicine_Frame = Frame(bg='#8FDFEA', bd=10, relief=RIDGE, padx=5, pady=10)
        self.Medicine_Name_Label = Label(self.Sell_Medicine_Frame, text="Medicine Name", font=('Calibri', 14, 'bold'),bg='#8FDFEA')
        self.Medicine_Name_Label.grid(row=0, column=0, padx=5)
        self.Medicine_Name_Entry = Entry(self.Sell_Medicine_Frame, bd=2, font=('Calibri', 15), relief=RIDGE,state="readonly", width=30)
        self.Medicine_Name_Entry.grid(row=0, column=1, pady=5)
        self.Company_Name_Label = Label(self.Sell_Medicine_Frame, text="Company Name", font=('Calibri', 14, 'bold'),bg='#8FDFEA')
        self.Company_Name_Label.grid(row=0, column=2, padx=5)
        self.Company_Name_Entry = Entry(self.Sell_Medicine_Frame, bd=2, font=('Calibri', 15), relief=RIDGE,state="readonly", width=30)
        self.Company_Name_Entry.grid(row=0, column=3, pady=5)
        self.Type_Of_Medicine_Label = Label(self.Sell_Medicine_Frame, text="Type Of Medicine",font=('Calibri', 14, 'bold'), bg='#8FDFEA')
        self.Type_Of_Medicine_Label.grid(row=0, column=4, padx=5)
        self.Type_Of_Medicine_Entry = Entry(self.Sell_Medicine_Frame, bd=2, font=('Calibri', 15), relief=RIDGE,state="readonly", width=27)
        self.Type_Of_Medicine_Entry.grid(row=0, column=5,pady=5)
        self.Usage_Label = Label(self.Sell_Medicine_Frame, text="Usage", font=('Calibri', 14, 'bold'), bg='#8FDFEA')
        self.Usage_Label.grid(row=1, column=0, padx=5)
        self.Usage_Entry = Entry(self.Sell_Medicine_Frame, bd=2, font=('Calibri', 15), relief=RIDGE, state="readonly",width=30)
        self.Usage_Entry.grid(row=1, column=1,pady=5)
        self.Mfg_Date_Label = Label(self.Sell_Medicine_Frame, text="Mfg. Date", font=('Calibri', 14, 'bold'),bg='#8FDFEA')
        self.Mfg_Date_Label.grid(row=1, column=2, padx=5)
        self.Mfg_Date_Entry = Entry(self.Sell_Medicine_Frame, bd=2, font=('Calibri', 15), relief=RIDGE,state="readonly", width=30)
        self.Mfg_Date_Entry.grid(row=1, column=3, pady=5)
        self.Expiry_Date_Label = Label(self.Sell_Medicine_Frame, text="Expiry Date", font=('Calibri', 14, 'bold'),bg='#8FDFEA')
        self.Expiry_Date_Label.grid(row=1, column=4, padx=5)
        self.Expiry_Date_Entry = Entry(self.Sell_Medicine_Frame, bd=2, font=('Calibri', 15), relief=RIDGE,state="readonly", width=27)
        self.Expiry_Date_Entry.grid(row=1, column=5, pady=5)
        self.Lot_Num_Label = Label(self.Sell_Medicine_Frame, text="Lot No.", font=('Calibri', 14, 'bold'), bg='#8FDFEA')
        self.Lot_Num_Label.grid(row=2, column=0, padx=5, pady=5)
        self.Lot_Num_Entry = Entry(self.Sell_Medicine_Frame, bd=2, font=('Calibri', 15), relief=RIDGE, state="readonly",width=30)
        self.Lot_Num_Entry.grid(row=2, column=1, pady=5)
        self.Price_Label = Label(self.Sell_Medicine_Frame, text="Price", font=('Calibri', 14, 'bold'), bg='#8FDFEA')
        self.Price_Label.grid(row=2, column=2, padx=5)
        self.Price_Entry = Entry(self.Sell_Medicine_Frame, bd=2, font=('Calibri', 15), relief=RIDGE, state="readonly",width=30)
        self.Price_Entry.grid(row=2, column=3, pady=5)
        self.Sell_Button = Button(self.Sell_Medicine_Frame, text="Sell Medicine", command=self.Sell_Button,font=('Calibri', 15), bg="green", fg='white', activebackground='green',activeforeground='white')
        self.Sell_Button.grid(row=2, column=4, columnspan=2, pady=5)
        self.Sell_Medicine_Frame.pack()
        self.style = Style()
        self.style.theme_use("default")
        self.style.configure("Treeview", background="#D3D3D3", rowheight=25, font=("Calibri", 12),fieldbackground='#D3D3D3')
        self.style.configure("Treeview.Heading", background="lightgreen", font=("Calibri", 12,'bold'))
        self.style.map("Treeview", background=[('selected', '#347083')])
        self.Show_All()
        self.Search_Entry.bind("<KeyRelease>", self.Enter_Search)
        self.Search_Table.bind("<<TreeviewSelect>>", self.Click_Medicine)
        self.Search_ComboBox.bind('<<ComboboxSelected>>', self.Type_Of_Medicine_Clicked)
    def Remove_Sell_Medicine(self):
        self.Main_Background.place_forget()
        self.Back_Add_Button.place_forget()
        self.Title_Label.pack_forget()
        self.Sell_Medicine_Frame.pack_forget()
        self.Results_Frame.pack_forget()
    def Back_Button(self):
        self.Remove_Sell_Medicine()
        Pharmacy()
    def Sell_Button(self):
        if not self.Search_Table.selection():
            messagebox.showerror("Error", "No Medicine Selected To Sell")
        else:
            cursor.execute(f"update Medicine set Sold=Sold+1,In_Storage=In_Storage-1 where Id={self.iid}")
            cursor.execute(f"update Medicine_Names set Total_Quantity=Total_Quantity-1 where Name='{self.Medicine_Name_Entry.get()}'")
            connect.commit()
            messagebox.showinfo("Sell Medicine","Medicine Sold Successfully")
            cursor.execute(f"select In_Storage from Medicine where In_Storage=10 and Id={self.iid}")
            res1=cursor.fetchone()
            if res1:
                messagebox.showwarning("Warning","There Are Only 10 Items Left From This Medicine")
            cursor.execute(f"select In_Storage from Medicine where In_Storage=1 and Id={self.iid}")
            res2 = cursor.fetchone()
            if res2:
                messagebox.showwarning("Warning", "There Is Only 1 Item Left From This Medicine")
            self.Show_All()
    def Search_Button(self):
        choice=''
        if self.Search_ComboBox.get()=='' and self.Search_Entry.get()=='':
            messagebox.showerror("Error","The Search By ComboBox And Search Entry Are Empty")
        elif self.Search_Entry.get() == '':
            messagebox.showerror("Error", "Enter What Do You Want To Search")
        else:
            if self.Search_ComboBox.get()=='Medicine Name':
                choice='Name'
            elif self.Search_ComboBox.get()=='Usage':
                choice='Usage'
            elif self.Search_ComboBox.get()=='Type Of Medicine':
                choice='Type'
            elif self.Search_ComboBox.get()=='Lot No.':
                choice='Lot_No'
            elif self.Search_ComboBox.get()=='Company Name':
                choice='Company_Name'
            cursor.execute(f"select Id,Name,Company_Name,Type,Usage,Mfg_Date,Expiry_Date,Lot_No,In_Storage,Price from Medicine where {choice}='{self.Search_Entry.get()}'")
            res = cursor.fetchall()
            for i in self.Search_Table.get_children():
                self.Search_Table.delete(i)
            for j in res:
                self.iid=j[0]
                self.Search_Table.insert('', END,iid=self.iid, values=j[1:])
    def Show_All(self):
        cursor.execute("select Id,Name,Company_Name,Type,Usage,Mfg_Date,Expiry_Date,Lot_No,In_Storage,Price from Medicine where In_Storage<>0")
        self.res = cursor.fetchall()
        for j in self.Search_Table.get_children():
            self.Search_Table.delete(j)
        for i in self.res:
            self.iid=i[0]
            self.Search_Table.insert('', END,iid=self.iid, values=i[1:])
    def Type_Of_Medicine_Clicked(self,e):
        self.Search_Entry.config(state="normal")
        if self.Search_ComboBox.get()=='Type Of Medicine':
            self.Type_Of_Medicine_Values = ('Tablets', 'Capsules', 'Liquids','Drops', 'Spray', 'Creams', 'Inhalers', 'Injections','Patches')
            self.Search_Entry.grid_forget()
            self.Search_Entry=Combobox(self.Search_Frame,font=('Calibri', 14),state="readonly",values=self.Type_Of_Medicine_Values)
            self.Search_Entry.grid(row=0,column=2)
        else:
            self.Search_Entry.grid_forget()
            self.Search_Entry = Entry(self.Search_Frame, bd=2, font=('Calibri', 15), relief=RIDGE)
            self.Search_Entry.grid(row=0, column=2)
            self.Search_Label=Label(self.Search_Entry, text="Search...", font=('Calibri', 13))
            self.Search_Label.place(x=10,y=0)
            self.Search_Entry.bind("<KeyRelease>", self.Enter_Search)
    def Enter_Search(self, e):
        if self.Search_Entry.get()=='':
            self.Search_Label.place(x=10, y=0)
            self.Show_All()
        else:
            self.Search_Label.place_forget()
            choice = ''
            if self.Search_ComboBox.get() == 'Medicine Name':
                choice = 'Name'
            elif self.Search_ComboBox.get() == 'Usage':
                choice = 'Usage'
            elif self.Search_ComboBox.get() == 'Type Of Medicine':
                choice = 'Type'
            elif self.Search_ComboBox.get() == 'Lot No.':
                choice = 'Lot_No'
            elif self.Search_ComboBox.get() == 'Company Name':
                choice = 'Company_Name'
            cursor.execute(f"select Id,Name,Company_Name,Type,Usage,Mfg_Date,Expiry_Date,Lot_No,In_Storage,Price from Medicine where In_Storage<>0 and {choice} like '%{self.Search_Entry.get()}%' ")
            self.res = cursor.fetchall()
            for j in self.Search_Table.get_children():
                self.Search_Table.delete(j)
            for i in self.res:
                self.iid = i[0]
                self.Search_Table.insert('', END, iid=self.iid, values=i[1:])
    def Click_Medicine(self,e):
        self.Medicine_Name_Entry.config(state="normal")
        self.Company_Name_Entry.config(state="normal")
        self.Usage_Entry.config(state="normal")
        self.Type_Of_Medicine_Entry.config(state="normal")
        self.Mfg_Date_Entry.config(state="normal")
        self.Expiry_Date_Entry.config(state="normal")
        self.Lot_Num_Entry.config(state="normal")
        self.Price_Entry.config(state="normal")
        self.Medicine_Name_Entry.delete(0, END)
        self.Company_Name_Entry.delete(0, END)
        self.Usage_Entry.delete(0, END)
        self.Type_Of_Medicine_Entry.delete(0, END)
        self.Mfg_Date_Entry.delete(0, END)
        self.Expiry_Date_Entry.delete(0, END)
        self.Lot_Num_Entry.delete(0, END)
        self.Price_Entry.delete(0, END)
        self.iid=self.Search_Table.focus()
        for i in self.Search_Table.selection():
            self.item = self.Search_Table.item(i)
            self.record = self.item['values']
        self.Medicine_Name_Entry.insert(END, self.record[0])
        self.Company_Name_Entry.insert(END, self.record[1])
        self.Type_Of_Medicine_Entry.insert(END, self.record[2])
        self.Usage_Entry.insert(END, self.record[3])
        self.Mfg_Date_Entry.insert(END,self.record[4])
        Expiry_Date = self.record[5]
        if Expiry_Date == '-':
            self.Expiry_Date_Label.grid_forget()
            self.Expiry_Date_Entry.grid_forget()
        else:
            self.Expiry_Date_Label.grid(row=1, column=4)
            self.Expiry_Date_Entry.grid(row=1, column=5)
            self.Expiry_Date_Entry.insert(END,self.record[5])
        self.Lot_Num_Entry.insert(END, self.record[6])
        self.Price_Entry.insert(END, self.record[8])
        self.Medicine_Name_Entry.config(state="readonly")
        self.Company_Name_Entry.config(state="readonly")
        self.Usage_Entry.config(state="readonly")
        self.Type_Of_Medicine_Entry.config(state="readonly")
        self.Mfg_Date_Entry.config(state="readonly")
        self.Expiry_Date_Entry.config(state="readonly")
        self.Lot_Num_Entry.config(state="readonly")
        self.Price_Entry.config(state="readonly")
class Show_Medicines:
    def __init__(self):
        self.Background_Image = ImageTk.PhotoImage(Image.open('medicine.jpg').resize((w, h)))
        self.Back_Image = ImageTk.PhotoImage(Image.open('back.png').resize((100, 100)))
        self.Main_Background = Label(image=self.Background_Image)
        self.Main_Background.place(x=0, y=0)
        self.Back_Add_Button = Button(image=self.Back_Image, command=self.Back_Button, bg='red', bd=0,activebackground='red')
        self.Back_Add_Button.place(x=0, y=0)
        self.Title_Label = Label(text="Show Medicines", bg='#7FFFD4', fg='#0020C2', font=('Calibri', 40, 'bold'), bd=10,relief=RIDGE)
        self.Title_Label.pack()
        self.Search_Frame=Frame(bg='#8FDFEA',bd=10,relief=RIDGE,padx=5,pady=10)
        self.Search_Frame.pack(pady=50)
        self.Search_By_Label=Label(self.Search_Frame,text="Search By",font=('Calibri', 14, 'bold'),bg='#8FDFEA')
        self.Search_By_Label.grid(row=0,column=0)
        self.Values=("Medicine Name","Company Name","Type Of Medicine","Usage","Lot No.")
        self.Search_ComboBox=Combobox(self.Search_Frame,font=('Calibri', 14),state="readonly",values=self.Values)
        self.Search_ComboBox.grid(row=0,column=1,padx=20)
        self.Search_Entry=Entry(self.Search_Frame,bd=2,font=('Calibri',15),relief=RIDGE,width=20,state="disabled")
        self.Search_Entry.grid(row=0,column=2)
        self.Search_Label=Label(self.Search_Entry, text="Search...", font=('Calibri', 13))
        self.Search_Label.place(x=10, y=0)
        self.Search_Button = Button(self.Search_Frame, text="Search",command=self.Search, font=('Calibri', 14), bg="green",fg='white', activebackground='green', activeforeground='white')
        self.Search_Button.grid(row=0, column=3,padx=20)
        self.Show_All_Button = Button(self.Search_Frame, text="Show All", command=self.Show_All, font=('Calibri', 14), bg="green", fg='white', activebackground='green', activeforeground='white')
        self.Show_All_Button.grid(row=0, column=4)
        self.Export_To_Excel_Button = Button(self.Search_Frame, text="Export To Excel", command=self.Export_To_Excel_Button, font=('Calibri', 14), bg="green", fg='white', activebackground='green', activeforeground='white')
        self.Export_To_Excel_Button.grid(row=0, column=5,padx=20)
        self.Update_Price_Button = Button(self.Search_Frame, text="Update Price", command=self.Update_Button, font=('Calibri', 14), bg="green", fg='white', activebackground='green', activeforeground='white')
        self.Update_Price_Button.grid(row=0, column=6)
        self.Update_Price_Entry =Entry(self.Search_Frame, bd=2, font=('Calibri', 15), relief=RIDGE, state='readonly')
        self.Update_Price_Entry.grid(row=0, column=7, padx=20)
        self.Delete_Button = Button(self.Search_Frame, text="Delete", command=self.Delete_Button, font=('Calibri', 14), bg="red",fg='white', activebackground='red', activeforeground='white')
        self.Delete_Button.grid(row=0, column=8)
        self.Table_Frame=Frame(bd=10,relief=RIDGE,bg='#8FDFEA')
        self.Table_Frame.pack()
        self.Columns = ("Medicine Name","Company Name","Type Of Medicine", "Usage", "Mfg. Date","Expiry Date", "Lot No.", "Quantity","Sold","In Storage", "Price")
        self.ScrollBar_Y = Scrollbar(self.Table_Frame)
        self.ScrollBar_Y.pack(side=RIGHT, fill=Y)
        self.Search_Table=Treeview(self.Table_Frame,columns=self.Columns,show='headings')
        self.ScrollBar_Y.config(command=self.Search_Table.yview)
        self.Search_Table.column("Medicine Name", width=160)
        self.Search_Table.column("Company Name", width=160)
        self.Search_Table.column("Mfg. Date", width=100,anchor=CENTER)
        self.Search_Table.column("Expiry Date", width=100,anchor=CENTER)
        self.Search_Table.column("Lot No.", width=120)
        self.Search_Table.column("Usage", width=200)
        self.Search_Table.column("Type Of Medicine", width=125,anchor=CENTER)
        self.Search_Table.column("Price", width=90,anchor=CENTER)
        self.Search_Table.column("Quantity", width=90,anchor=CENTER)
        self.Search_Table.column("Sold", width=90,anchor=CENTER)
        self.Search_Table.column("In Storage", width=90,anchor=CENTER)
        self.Search_Table.heading("Medicine Name", text="Medicine Name",anchor=CENTER)
        self.Search_Table.heading("Company Name", text="Company Name",anchor=CENTER)
        self.Search_Table.heading("Mfg. Date", text="Mfg. Date",anchor=CENTER)
        self.Search_Table.heading("Expiry Date", text="Expiry Date",anchor=CENTER)
        self.Search_Table.heading("Lot No.", text="Lot No.",anchor=CENTER)
        self.Search_Table.heading("Type Of Medicine", text="Type Of Medicine",anchor=CENTER)
        self.Search_Table.heading("Usage", text="Usage",anchor=CENTER)
        self.Search_Table.heading("Price", text="Price",anchor=CENTER)
        self.Search_Table.heading("Quantity", text="Quantity",anchor=CENTER)
        self.Search_Table.heading("Sold", text="Sold",anchor=CENTER)
        self.Search_Table.heading("In Storage", text="In Storage",anchor=CENTER)
        self.Search_Table.pack()
        self.style = Style()
        self.style.theme_use("default")
        self.style.configure("Treeview",background="#D3D3D3",rowheight=25,font=("Calibri",12),fieldbackground='#D3D3D3')
        self.style.configure("Treeview.Heading",background="lightgreen",font=("Calibri",12,'bold'))
        self.style.map("Treeview",background=[('selected', '#347083')])
        self.Show_All()
        self.Search_Entry.bind("<KeyRelease>", self.Enter_Search)
        self.Search_Table.bind("<<TreeviewSelect>>", self.Click_Medicine)
        self.Search_ComboBox.bind('<<ComboboxSelected>>', self.Type_Of_Medicine_Clicked)
    def Remove_Show_Medicines(self):
        self.Main_Background.place_forget()
        self.Back_Add_Button.place_forget()
        self.Title_Label.pack_forget()
        self.Search_Frame.pack_forget()
        self.Table_Frame.pack_forget()
    def Back_Button(self):
        self.Remove_Show_Medicines()
        Pharmacy()
    def Show_All(self):
        cursor.execute("select * from Medicine")
        self.res = cursor.fetchall()
        for j in self.Search_Table.get_children():
            self.Search_Table.delete(j)
        for i in self.res:
            self.iid=i[0]
            self.Search_Table.insert('', END,iid=self.iid, values=i[1:])
    def Search(self):
        choice=''
        if self.Search_ComboBox.get()=='' and self.Search_Entry.get()=='':
            messagebox.showerror("Error","The Search By ComboBox And Search Entry Are Empty")
        elif self.Search_Entry.get()=='':
            messagebox.showerror("Error","Enter What Do You Want To Search")
        else:
            if self.Search_ComboBox.get()=='Medicine Name':
                choice='Name'
            elif self.Search_ComboBox.get()=='Usage':
                choice='Usage'
            elif self.Search_ComboBox.get()=='Type Of Medicine':
                choice='Type'
            elif self.Search_ComboBox.get()=='Lot No.':
                choice='Lot_No'
            elif self.Search_ComboBox.get()=='Company Name':
                choice='Company_Name'
            cursor.execute(f"select * from Medicine where {choice}='{self.Search_Entry.get()}'")
            res = cursor.fetchall()
            for i in self.Search_Table.get_children():
                self.Search_Table.delete(i)
            for j in res:
                self.Search_Table.insert('', END, values=j[1:])
    def Export_To_Excel_Button(self):
        from openpyxl import Workbook
        from openpyxl.styles import Side, Font, Border, Alignment,PatternFill
        from openpyxl.utils import coordinate_to_tuple, get_column_letter
        from tkinter import filedialog
        wb =Workbook()
        sheet = wb.active
        sheet.merge_cells('E2:F2')
        today=sheet['E2']
        date=datetime.date.today().strftime("%d/%m/%Y")
        today.value=date
        today.font = Font(bold=True, size=16)
        today.alignment = Alignment(horizontal='center', vertical='center')
        today_text = sheet['D2']
        today_text.value = "Date"
        today_text.font = Font(bold=True, size=16)
        today_text.alignment = Alignment(horizontal='center', vertical='center')
        start_cell = 'B4'
        self.Columns = ("Medicine Name","Company Name","Type Of Medicine", "Usage", "Mfg. Date","Expiry Date", "Lot No.", "Quantity","Sold","In Storage", "Price"),
        start_row, start_col = coordinate_to_tuple(start_cell)
        for row_index, row_data in enumerate(self.Columns, start=start_row):
            for col_index, value in enumerate(row_data, start=start_col):
                cell = sheet.cell(row=row_index, column=col_index, value=value)
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.fill=PatternFill(start_color="0000FF00",fill_type="solid")
        additional_start_row = start_row + 1
        records=[]
        for rec in self.res:
            x = list(rec)
            x.remove(rec[0])
            records.append(x)
        for row_index, row_data in enumerate(records, start=additional_start_row):
            for col_index, value in enumerate(row_data, start=start_col):
                cell = sheet.cell(row=row_index, column=col_index, value=value)
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.fill=PatternFill(start_color="00CCCCFF",fill_type="solid")
        for i, row in enumerate(sheet.iter_rows(min_row=start_row, min_col=start_col), start=start_row):
            for j, cell in enumerate(row, start=start_col):
                cell.border = Border(
                    left=Side(border_style="thin", color="000000"),
                    right=Side(border_style="thin", color="000000"),
                    top=Side(border_style="thin", color="000000"),
                    bottom=Side(border_style="thin", color="000000")
                )
                cell.font = Font(size=14)
                if i == 4:
                    cell.font = Font(bold=True, size=16)
        max_col = sheet.max_column
        for col in range(start_col, max_col + 1):
            column_letter = get_column_letter(col)
            if column_letter=='E':
                sheet.column_dimensions[column_letter].width = 40
            else:
                sheet.column_dimensions[column_letter].width = 25
        file_path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel files", "*.xlsx")])
        if file_path:
            wb.save(file_path)
            messagebox.showinfo("Save File","Excel File Saved Successfully")
        else:
            messagebox.showerror("Error","Excel File Is Not Saved ")
    def Update_Button(self):
        if not self.Search_Table.selection():
            messagebox.showerror("Error", "No Medicine Selected To Update Price")
        elif not self.Update_Price_Entry.get().isdigit():
            messagebox.showerror("Error","Wrong Entry In Price")
        else:
            cursor.execute(f"""
                update Medicine set Price='{int(self.Update_Price_Entry.get())}'where Id={self.iid}
                """)
            connect.commit()
            messagebox.showinfo("Medicine Update", "Medicine Updated Successfully")
            self.Show_All()
            self.Update_Price_Entry.delete(0, END)
            self.Update_Price_Entry.config(state="readonly")
    def Delete_Button(self):
        if not self.Search_Table.selection():
            messagebox.showerror("Error","No Medicine Selected To Delete")
        else:
            answer=messagebox.askyesno("Delete","Are You Sure You Want To Delete This Medicine")
            if answer:
                cursor.execute(f"select In_Storage,Name from Medicine where Id={self.iid}")
                r=cursor.fetchone()
                cursor.execute(f"update Medicine_Names set Total_Quantity=Total_Quantity-{r[0]} where Name='{r[1]}'")
                cursor.execute(f"delete from Medicine where Id={self.iid}")
                connect.commit()
                messagebox.showinfo("Medicine Deleted","Medicine Deleted Successfully")
                self.Show_All()
        self.Update_Price_Entry.delete(0, END)
        self.Update_Price_Entry.config(state="readonly")
    def Type_Of_Medicine_Clicked(self,e):
        self.Search_Entry.config(state="normal")
        if self.Search_ComboBox.get()=='Type Of Medicine':
            self.Type_Of_Medicine_Values = ('Tablets', 'Capsules', 'Liquids','Drops', 'Spray', 'Creams', 'Inhalers', 'Injections','Patches')
            self.Search_Entry.grid_forget()
            self.Search_Entry=Combobox(self.Search_Frame,font=('Calibri', 14),state="readonly",values=self.Type_Of_Medicine_Values)
            self.Search_Entry.grid(row=0,column=2)
        else:
            self.Search_Entry.grid_forget()
            self.Search_Entry = Entry(self.Search_Frame, bd=2, font=('Calibri', 15), relief=RIDGE)
            self.Search_Entry.grid(row=0, column=2)
            self.Search_Label=Label(self.Search_Entry, text="Search...", font=('Calibri', 13))
            self.Search_Label.place(x=10,y=0)
            self.Search_Entry.bind("<KeyRelease>", self.Enter_Search)
    def Click_Medicine(self,e):
        self.Update_Price_Entry.config(state="normal")
        self.iid=self.Search_Table.focus()
        self.Update_Price_Entry.delete(0, END)
        for i in self.Search_Table.selection():
            self.item = self.Search_Table.item(i)
            self.record = self.item['values']
        self.Update_Price_Entry.insert(END, self.record[10])
    def Enter_Search(self, e):
        if self.Search_Entry.get()=='':
            self.Search_Label.place(x=10, y=0)
            self.Show_All()
        else:
            self.Search_Label.place_forget()
            choice = ''
            if self.Search_ComboBox.get()=='Medicine Name':
                choice='Name'
            elif self.Search_ComboBox.get()=='Usage':
                choice='Usage'
            elif self.Search_ComboBox.get()=='Type Of Medicine':
                choice='Type'
            elif self.Search_ComboBox.get()=='Lot No.':
                choice='Lot_No'
            elif self.Search_ComboBox.get()=='Company Name':
                choice='Company_Name'
            cursor.execute(f"select * from Medicine where {choice} like '%{self.Search_Entry.get()}%' ")
            self.res = cursor.fetchall()
            for j in self.Search_Table.get_children():
                self.Search_Table.delete(j)
            for i in self.res:
                self.iid = i[0]
                self.Search_Table.insert('', END, iid=self.iid, values=i[1:])
class Show_Medicines_Quantity:
    def __init__(self):
        self.Background_Image = ImageTk.PhotoImage(Image.open('medicine.jpg').resize((w, h)))
        self.Back_Image = ImageTk.PhotoImage(Image.open('back.png').resize((100, 100)))
        self.Main_Background = Label(image=self.Background_Image)
        self.Main_Background.place(x=0, y=0)
        self.Back_Add_Button = Button(image=self.Back_Image, command=self.Back_Button, bg='red', bd=0,activebackground='red')
        self.Back_Add_Button.place(x=0, y=0)
        self.Title_Label = Label(text="Show Medicines Quantity", bg='#7FFFD4', fg='#0020C2', font=('Calibri', 40, 'bold'), bd=10,relief=RIDGE)
        self.Title_Label.pack()
        self.Search_Frame=Frame(bg='#8FDFEA',bd=10,relief=RIDGE,padx=20,pady=10)
        self.Search_Frame.pack(pady=50)
        self.Search_Name_Label=Label(self.Search_Frame, text="Search Name", font=('Calibri', 14, 'bold'), bg='#8FDFEA')
        self.Search_Name_Label.grid(row=0, column=0)
        self.Search_Entry=Entry(self.Search_Frame,bd=2,font=('Calibri',15),relief=RIDGE)
        self.Search_Entry.grid(row=0,column=1,padx=20)
        self.Search_Label = Label(self.Search_Entry, text="Search...", font=('Calibri', 13))
        self.Search_Label.place(x=10, y=0)
        self.Search_Button = Button(self.Search_Frame, text="Search",command=self.Search, font=('Calibri', 14), bg="green",fg='white', activebackground='green', activeforeground='white')
        self.Search_Button.grid(row=0, column=2)
        self.Show_All_Button = Button(self.Search_Frame, text="Show All", command=self.Show_All, font=('Calibri', 14), bg="green", fg='white', activebackground='green', activeforeground='white')
        self.Show_All_Button.grid(row=0, column=3,padx=20)
        self.Table_Frame=Frame(bd=10,relief=RIDGE,bg='#8FDFEA')
        self.Table_Frame.pack()
        self.Columns = ("Medicine Name", "Total Quantity")
        self.ScrollBar_Y = Scrollbar(self.Table_Frame)
        self.ScrollBar_Y.pack(side=RIGHT, fill=Y)
        self.Search_Table=Treeview(self.Table_Frame,columns=self.Columns,show='headings')
        self.ScrollBar_Y.config(command=self.Search_Table.yview)
        self.Search_Table.column("Medicine Name", width=170)
        self.Search_Table.column("Total Quantity", width=120,anchor=CENTER)
        self.Search_Table.heading("Medicine Name", text="Medicine Name",anchor=CENTER)
        self.Search_Table.heading("Total Quantity", text="Total Quantity",anchor=CENTER)
        self.Search_Table.pack()
        self.style = Style()
        self.style.theme_use("default")
        self.style.configure("Treeview",background="#D3D3D3",rowheight=25,font=("Calibri",12),fieldbackground='#D3D3D3')
        self.style.configure("Treeview.Heading",background="lightgreen",font=("Calibri",12,'bold'))
        self.style.map("Treeview",background=[('selected', '#347083')])
        self.Show_All()
        self.Search_Entry.bind("<KeyRelease>", self.Enter_Search)
    def Remove_Show_Medicines_Quantity(self):
        self.Main_Background.place_forget()
        self.Back_Add_Button.place_forget()
        self.Title_Label.pack_forget()
        self.Search_Frame.pack_forget()
        self.Table_Frame.pack_forget()
    def Back_Button(self):
        self.Remove_Show_Medicines_Quantity()
        Pharmacy()
    def Show_All(self):
        cursor.execute("select * from Medicine_Names")
        self.res = cursor.fetchall()
        for j in self.Search_Table.get_children():
            self.Search_Table.delete(j)
        for i in self.res:
            self.iid=i[0]
            self.Search_Table.insert('', END,iid=self.iid, values=i[1:])
    def Search(self):
        if self.Search_Entry.get()=='':
            messagebox.showerror("Error","Name Is Empty")
        else:
            cursor.execute(f"select * from Medicine_Names where Name='{self.Search_Entry.get()}'")
            res = cursor.fetchall()
            for i in self.Search_Table.get_children():
                self.Search_Table.delete(i)
            for j in res:
                self.Search_Table.insert('', END, values=j[1:])
    def Enter_Search(self, e):
        if self.Search_Entry.get()=='':
            self.Search_Label.place(x=10, y=0)
            self.Show_All()
        else:
            self.Search_Label.place_forget()
            cursor.execute(f"select * from Medicine_Names where Name like '%{self.Search_Entry.get()}%' ")
            self.res = cursor.fetchall()
            for j in self.Search_Table.get_children():
                self.Search_Table.delete(j)
            for i in self.res:
                self.iid = i[0]
                self.Search_Table.insert('', END, iid=self.iid, values=i[1:])
class Show_Expired_Medicines:
    def __init__(self):
        self.Background_Image = ImageTk.PhotoImage(Image.open('medicine.jpg').resize((w, h)))
        self.Back_Image = ImageTk.PhotoImage(Image.open('back.png').resize((100, 100)))
        self.Main_Background = Label(image=self.Background_Image)
        self.Main_Background.place(x=0, y=0)
        self.Back_Add_Button = Button(image=self.Back_Image, command=self.Back_Button, bg='red', bd=0,activebackground='red')
        self.Back_Add_Button.place(x=0, y=0)
        self.Title_Label = Label(text="Show Expired Medicines", bg='#7FFFD4', fg='#0020C2', font=('Calibri', 40, 'bold'), bd=10,relief=RIDGE)
        self.Title_Label.pack()
        self.Table_Frame = Frame(bd=10, relief=RIDGE, bg='#8FDFEA')
        self.Table_Frame.pack(pady=50)
        self.Columns = ("Medicine Name","Company Name","Type Of Medicine", "Usage", "Mfg. Date","Expiry Date", "Lot No.","Sold","In Storage", "Price")
        self.ScrollBar_Y = Scrollbar(self.Table_Frame)
        self.ScrollBar_Y.pack(side=RIGHT, fill=Y)
        self.Search_Table = Treeview(self.Table_Frame, columns=self.Columns, show='headings')
        self.ScrollBar_Y.config(command=self.Search_Table.yview)
        self.Search_Table.column("Medicine Name", width=160)
        self.Search_Table.column("Company Name", width=160)
        self.Search_Table.column("Mfg. Date", width=100, anchor=CENTER)
        self.Search_Table.column("Expiry Date", width=100, anchor=CENTER)
        self.Search_Table.column("Lot No.", width=120)
        self.Search_Table.column("Usage", width=200)
        self.Search_Table.column("Type Of Medicine", width=130, anchor=CENTER)
        self.Search_Table.column("Price", width=90, anchor=CENTER)
        self.Search_Table.column("Sold", width=90, anchor=CENTER)
        self.Search_Table.column("In Storage", width=90, anchor=CENTER)
        self.Search_Table.heading("Medicine Name", text="Medicine Name", anchor=CENTER)
        self.Search_Table.heading("Company Name", text="Company Name", anchor=CENTER)
        self.Search_Table.heading("Mfg. Date", text="Mfg. Date", anchor=CENTER)
        self.Search_Table.heading("Expiry Date", text="Expiry Date", anchor=CENTER)
        self.Search_Table.heading("Lot No.", text="Lot No.", anchor=CENTER)
        self.Search_Table.heading("Type Of Medicine", text="Type Of Medicine", anchor=CENTER)
        self.Search_Table.heading("Usage", text="Usage", anchor=CENTER)
        self.Search_Table.heading("Price", text="Price", anchor=CENTER)
        self.Search_Table.heading("Sold", text="Sold", anchor=CENTER)
        self.Search_Table.heading("In Storage", text="In Storage", anchor=CENTER)
        self.Search_Table.pack()
        self.style = Style()
        self.style.theme_use("default")
        self.style.configure("Treeview", background="#D3D3D3", rowheight=25, font=("Calibri", 12),fieldbackground='#D3D3D3')
        self.style.configure("Treeview.Heading", background="lightgreen", font=("Calibri", 12,'bold'))
        self.style.map("Treeview", background=[('selected', '#347083')])
        self.Expired_Medicines()
    def Expired_Medicines(self):
        cursor.execute("select Expiry_Date from Medicine ")
        expiry = cursor.fetchall()
        for e in expiry:
            date = e[0]
            if date=='-':
                continue
            remove_slashes = date.split('/')
            result = f"{remove_slashes[2]}-{remove_slashes[1]}-{remove_slashes[0]}"
            cursor.execute(f"update Medicine set Expiry_Date='{result}' where Expiry_Date='{date}'")
            connect.commit()
        cursor.execute(f"select * from Medicine where Expiry_Date<='{datetime.date.today()}'")
        res = cursor.fetchall()
        for i in self.Search_Table.get_children():
            self.Search_Table.delete(i)
        for j in res:
            x=list(j)
            x.remove(x[8])
            date = x[6]
            if date=='-':
                continue
            remove_dashes = date.split('-')
            result = f"{remove_dashes[2]}/{remove_dashes[1]}/{remove_dashes[0]}"
            x[6] = result
            self.Search_Table.insert('', END, values=x[1:])
        cursor.execute("select Expiry_Date from Medicine ")
        expiry2 = cursor.fetchall()
        for e in expiry2:
            date = e[0]
            if date=='-':
                continue
            remove_dashes = date.split('-')
            result = f"{remove_dashes[2]}/{remove_dashes[1]}/{remove_dashes[0]}"
            cursor.execute(f"update Medicine set Expiry_Date='{result}' where Expiry_Date='{date}'")
            connect.commit()
    def Remove_Show_Expired_Medicines(self):
        self.Main_Background.place_forget()
        self.Back_Add_Button.place_forget()
        self.Title_Label.pack_forget()
        self.Table_Frame.pack_forget()
    def Back_Button(self):
        self.Remove_Show_Expired_Medicines()
        Pharmacy()
class Show_Missing_Medicines:
    def __init__(self):
        self.Background_Image = ImageTk.PhotoImage(Image.open('medicine.jpg').resize((w, h)))
        self.Back_Image = ImageTk.PhotoImage(Image.open('back.png').resize((100, 100)))
        self.Main_Background = Label(image=self.Background_Image)
        self.Main_Background.place(x=0, y=0)
        self.Back_Add_Button = Button(image=self.Back_Image, command=self.Back_Button, bg='red', bd=0,activebackground='red')
        self.Back_Add_Button.place(x=0, y=0)
        self.Title_Label = Label(text="Show Missing Medicines", bg='#7FFFD4', fg='#0020C2', font=('Calibri', 40, 'bold'), bd=10,relief=RIDGE)
        self.Title_Label.pack()
        self.Table_Frame = Frame(bd=10, relief=RIDGE, bg='#8FDFEA')
        self.Table_Frame.pack(pady=50)
        self.Columns = ("Medicine Name",)
        self.ScrollBar_Y = Scrollbar(self.Table_Frame)
        self.ScrollBar_Y.pack(side=RIGHT, fill=Y)
        self.Search_Table = Treeview(self.Table_Frame, columns=self.Columns, show='headings')
        self.ScrollBar_Y.config(command=self.Search_Table.yview)
        self.Search_Table.column("Medicine Name", width=170)
        self.Search_Table.heading("Medicine Name", text="Medicine Name", anchor=CENTER)
        self.Search_Table.pack()
        self.style = Style()
        self.style.theme_use("default")
        self.style.configure("Treeview", background="#D3D3D3", rowheight=25, font=("Calibri", 12),fieldbackground='#D3D3D3')
        self.style.configure("Treeview.Heading", background="lightgreen", font=("Calibri", 12,'bold'))
        self.style.map("Treeview", background=[('selected', '#347083')])
        cursor.execute("select Name from Medicine_Names where Total_Quantity=0")
        res = cursor.fetchall()
        for i in self.Search_Table.get_children():
            self.Search_Table.delete(i)
        for j in res:
            self.Search_Table.insert('', END, values=j)
    def Remove_Show_Missing_Medicines(self):
        self.Main_Background.place_forget()
        self.Back_Add_Button.place_forget()
        self.Title_Label.pack_forget()
        self.Table_Frame.pack_forget()
    def Back_Button(self):
        self.Remove_Show_Missing_Medicines()
        Pharmacy()
g=Tk()
w=g.winfo_screenwidth()
h=g.winfo_screenheight()
g.title("Pharmacy Management System")
g.state('zoomed')
Pharmacy()
g.mainloop()
